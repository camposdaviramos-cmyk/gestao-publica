import csv
import io
import json
from html import escape
from flask import request, send_file, jsonify
from auth import require, ApiError
from db import get_db, settings, audit
from domain import MODULES, now

def install_reports(app):
 @app.get('/api/reports/<module>')
 def report(module):
  if module not in MODULES: raise ApiError('Relatório não encontrado.',404)
  require(module); fmt=request.args.get('format','json'); cfg=settings()
  if fmt not in ['json','csv','xlsx','pdf','docx','html']: raise ApiError('Formato não suportado.')
  from integration_signatures import authorize_report,report_response
  from erp_core import entity_access
  entity=1
  if fmt not in ['json','html'] and (cfg['signature_required'] or module in cfg['signature_reports']):entity=entity_access(request.args.get('entity',1))
  authorize_report(module,entity,fmt)
  where='module=?'; args=[module]
  for param,col in [('from','created_at'),('to','created_at')]:
   if request.args.get(param):
    from datetime import datetime
    try: datetime.strptime(request.args[param],'%Y-%m-%d')
    except ValueError: raise ApiError('Período inválido.')
    from datetime import timedelta, timezone
    from zoneinfo import ZoneInfo
    boundary=datetime.strptime(request.args[param],'%Y-%m-%d').replace(tzinfo=ZoneInfo('America/Sao_Paulo'))
    if param=='to': boundary+=timedelta(days=1)
    where+=f' AND {col} '+('>=' if param=='from' else '<')+' ?'; args.append(boundary.astimezone(timezone.utc).isoformat(timespec='seconds'))
  if request.args.get('from') and request.args.get('to') and request.args['from']>request.args['to']: raise ApiError('O início deve ser anterior ao fim do período.')
  db=get_db(); count=db.execute('SELECT count(*) FROM records WHERE '+where,args).fetchone()[0]
  if count>10000: raise ApiError('Selecione um período menor: limite de 10.000 registros por relatório.')
  records=db.execute('SELECT * FROM records WHERE '+where+' ORDER BY id',args).fetchall()
  headers=['ID','Título','Unidade','Situação','Valor (R$)','Atualizado em']
  rows=[[r['id'],r['title'],r['department'],r['status'],f"{r['amount']/100:.2f}",r['updated_at']] for r in records]
  audit('Relatório gerado',module,detail={'format':fmt,'rows':len(rows)})
  page = max(1, int(request.args.get('page', 1)))
  per_page = max(1, min(500, int(request.args.get('per_page', 50))))
  total_count = len(rows)
  total_pages = (total_count + per_page - 1) // per_page if per_page > 0 else 1
  start_idx = (page - 1) * per_page
  paginated_rows = rows[start_idx:start_idx + per_page]

  if fmt=='json':
   return jsonify(headers=headers,rows=paginated_rows,all_rows=rows,page=page,per_page=per_page,total_pages=total_pages,total_records=total_count,total=sum(r['amount'] for r in records)/100,count=len(paginated_rows),generated_at=now())
  if fmt=='html':
   table_rows = ''.join('<tr>' + ''.join(f'<td>{escape(str(c))}</td>' for c in r) + '</tr>' for r in paginated_rows)
   label = MODULES[module]['label']
   html_content = f"""<!doctype html>
<html lang="pt-BR">
<head>
<meta charset="utf-8">
<title>{escape(cfg['municipality'])} · {escape(label)}</title>
<style>
body {{ font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif; margin: 2rem; color: #1e293b; }}
.header {{ display: flex; justify-content: space-between; align-items: center; border-bottom: 2px solid #0284c7; padding-bottom: 1rem; margin-bottom: 1.5rem; }}
h1 {{ margin: 0; font-size: 1.5rem; color: #0284c7; }}
.meta {{ font-size: 0.85rem; color: #64748b; }}
table {{ width: 100%; border-collapse: collapse; margin-top: 1rem; font-size: 0.9rem; }}
th, td {{ border: 1px solid #cbd5e1; padding: 0.6rem 0.8rem; text-align: left; }}
th {{ background: #f1f5f9; color: #334155; font-weight: 600; }}
tr:nth-child(even) {{ background: #f8fafc; }}
.pagination-info {{ margin-top: 1.5rem; font-size: 0.85rem; color: #64748b; display: flex; justify-content: space-between; }}
.btn-print {{ background: #0284c7; color: white; border: none; padding: 0.5rem 1rem; font-size: 0.9rem; font-weight: 600; border-radius: 6px; cursor: pointer; }}
@media print {{
  .no-print {{ display: none !important; }}
  body {{ margin: 0; }}
  table {{ font-size: 8pt; }}
}}
</style>
</head>
<body>
<div class="header">
  <div>
    <h1>{escape(cfg['municipality'])} · {escape(label)}</h1>
    <div class="meta">Emitido em {now()} | Relatório Oficial Administrativo</div>
  </div>
  <div class="no-print">
    <button class="btn-print" onclick="window.print()">🖨️ Imprimir / Salvar PDF</button>
  </div>
</div>
<table>
  <thead><tr>{''.join(f'<th>{escape(h)}</th>' for h in headers)}</tr></thead>
  <tbody>{table_rows}</tbody>
</table>
<div class="pagination-info">
  <span>Página {page} de {total_pages} (Exibindo {len(paginated_rows)} de {total_count} registros)</span>
  <span>Total consolidado: R$ {sum(r['amount'] for r in records)/100:,.2f}</span>
</div>
</body>
</html>"""
   return html_content, 200, {'Content-Type': 'text/html; charset=utf-8'}
  def safe(value):
   v=str(value)
   return "'"+v if v.lstrip().startswith(('=','+','-','@','\t','\r')) else v
  buffer=io.BytesIO(); label=MODULES[module]['label']; title=f"{cfg['municipality']} · {label}"
  if fmt=='csv':
   stream=io.StringIO(); writer=csv.writer(stream,delimiter=';'); writer.writerow(headers); writer.writerows([[safe(v) for v in r] for r in rows]); buffer.write(stream.getvalue().encode('utf-8-sig')); mime='text/csv'
  elif fmt=='xlsx':
   from openpyxl import Workbook
   from openpyxl.styles import Font, PatternFill
   wb=Workbook(); ws=wb.active; ws.title='Relatório'; ws.append(headers)
   for row in rows: ws.append([safe(v) for v in row])
   for cell in ws[1]: cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='176B54')
   ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
   for column,width in zip('ABCDEF',[10,55,35,22,20,30]): ws.column_dimensions[column].width=width
   wb.save(buffer); mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
  elif fmt=='docx':
   from docx import Document
   doc=Document(); doc.add_heading(title,0); doc.add_paragraph('Emitido em '+now()); table=doc.add_table(rows=1,cols=len(headers)); table.style='Light Shading Accent 1'
   for cell,text in zip(table.rows[0].cells,headers): cell.text=text
   for row in rows:
    for cell,value in zip(table.add_row().cells,row): cell.text=str(value)
   doc.save(buffer); mime='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
  else:
   from reportlab.lib import colors
   from reportlab.lib.pagesizes import A4, landscape
   from reportlab.lib.styles import getSampleStyleSheet
   from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, LongTable, TableStyle
   styles=getSampleStyleSheet(); body=styles['BodyText']; body.fontSize=8; body.leading=11
   table=LongTable([[Paragraph(escape(str(v)),body) for v in row] for row in [headers]+rows],colWidths=[35,215,150,95,80,180],repeatRows=1)
   table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#ddf2e9')),('VALIGN',(0,0),(-1,-1),'TOP'),('BOTTOMPADDING',(0,0),(-1,-1),9),('TOPPADDING',(0,0),(-1,-1),9),('LINEBELOW',(0,0),(-1,-1),.3,colors.HexColor('#dfe5e2'))]))
   def footer(canvas,doc):
    canvas.setFont('Helvetica',8); canvas.drawString(24,18,'Rio Gestão · Relatório administrativo'); canvas.drawRightString(817,18,f'Página {doc.page}')
   SimpleDocTemplate(buffer,pagesize=landscape(A4),leftMargin=24,rightMargin=24).build([Paragraph(escape(title),styles['Title']),Paragraph('Emitido em '+now(),body),Spacer(1,16),table],onFirstPage=footer,onLaterPages=footer); mime='application/pdf'
  return report_response(buffer,mime,f'rio-{module}-{now()[:10]}.{fmt}',module,entity)
